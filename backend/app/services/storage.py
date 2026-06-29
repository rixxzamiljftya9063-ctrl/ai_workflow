from pathlib import Path
from uuid import uuid4

from fastapi import UploadFile

from app.config import get_settings


settings = get_settings()


def project_root(project_id: int) -> Path:
    path = settings.storage_path / "projects" / str(project_id)
    path.mkdir(parents=True, exist_ok=True)
    return path


def files_dir(project_id: int) -> Path:
    path = project_root(project_id) / "files"
    path.mkdir(parents=True, exist_ok=True)
    return path


def workspace_dir(project_id: int) -> Path:
    path = project_root(project_id) / "workspace"
    path.mkdir(parents=True, exist_ok=True)
    return path


def outputs_dir(project_id: int) -> Path:
    path = project_root(project_id) / "outputs"
    path.mkdir(parents=True, exist_ok=True)
    return path


async def save_upload(project_id: int, upload: UploadFile) -> tuple[str, Path, bytes]:
    suffix = Path(upload.filename or "upload").suffix
    filename = f"{uuid4().hex}{suffix}"
    target = files_dir(project_id) / filename
    content = await upload.read()
    target.write_bytes(content)
    return filename, target, content


def extract_text(path: Path, content: bytes, mime_type: str = "") -> str:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md", ".csv", ".json", ".py"} or mime_type.startswith("text/"):
        for encoding in ("utf-8", "utf-8-sig", "gb18030", "latin-1"):
            try:
                return content.decode(encoding)
            except UnicodeDecodeError:
                continue
        return content.decode("utf-8", errors="ignore")
    if suffix == ".pdf":
        try:
            from pypdf import PdfReader

            reader = PdfReader(str(path))
            return "\n".join(page.extract_text() or "" for page in reader.pages)
        except Exception as exc:
            return f"[PDF text extraction failed: {exc}]"
    if suffix == ".docx":
        try:
            from docx import Document

            doc = Document(str(path))
            return "\n".join(paragraph.text for paragraph in doc.paragraphs)
        except Exception as exc:
            return f"[DOCX text extraction failed: {exc}]"
    if suffix == ".xlsx":
        try:
            from openpyxl import load_workbook

            workbook = load_workbook(str(path), read_only=True, data_only=True)
            lines: list[str] = []
            for sheet in workbook.worksheets:
                lines.append(f"# {sheet.title}")
                for row in sheet.iter_rows(values_only=True):
                    values = [str(cell) for cell in row if cell is not None]
                    if values:
                        lines.append("\t".join(values))
            return "\n".join(lines)
        except Exception as exc:
            return f"[XLSX text extraction failed: {exc}]"
    return ""
